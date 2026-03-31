"""Excel export for parsed store data."""

from __future__ import annotations

import logging
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from core.diff import (
    CHANGE_SHEET_COLUMNS,
    DiffResult,
    build_snapshot_rows,
    compute_diff,
    load_snapshot_with_meta,
    save_snapshot,
)
from core.models import STORE_OUTPUT_COLUMNS, StoreRecord

DATA_COLUMNS = STORE_OUTPUT_COLUMNS
STATS_COLUMNS = ["network", "stores_count"]

DATA_SHEET_NAME = "Актуальные данные"
CHANGES_SHEET_NAME = "Изменения"
STATS_SHEET_NAME = "Статистика"
MAX_COLUMN_WIDTH = 80

logger = logging.getLogger(__name__)


def _build_data_df(stores: list[StoreRecord]) -> pd.DataFrame:
    rows = [{column: row.get(column) for column in DATA_COLUMNS} for row in build_snapshot_rows(stores)]
    data_df = pd.DataFrame(rows, columns=DATA_COLUMNS)
    if data_df.empty:
        return pd.DataFrame(columns=DATA_COLUMNS)
    return data_df.sort_values(
        by=["network", "region", "city", "address", "source_url"],
        kind="stable",
        na_position="last",
    ).reset_index(drop=True)


def _build_changes_df(diff_result: DiffResult) -> pd.DataFrame:
    rows = diff_result.to_rows()
    if not rows:
        return pd.DataFrame(columns=CHANGE_SHEET_COLUMNS)
    return pd.DataFrame(rows, columns=CHANGE_SHEET_COLUMNS)


def _build_stats_df(data_df: pd.DataFrame) -> pd.DataFrame:
    if data_df.empty:
        return pd.DataFrame(columns=STATS_COLUMNS)
    return (
        data_df.groupby("network", dropna=False)
        .size()
        .reset_index(name="stores_count")
        .sort_values(by=["stores_count", "network"], ascending=[False, True], kind="stable")
        .reset_index(drop=True)
    )


def _apply_worksheet_formatting(worksheet) -> None:  # type: ignore[no-untyped-def]
    worksheet.freeze_panes = "A2"
    if worksheet.max_row >= 1 and worksheet.max_column >= 1:
        worksheet.auto_filter.ref = worksheet.dimensions

    for cell in worksheet[1]:
        cell.font = Font(bold=True)

    for column_index in range(1, worksheet.max_column + 1):
        letter = get_column_letter(column_index)
        max_length = 0
        for cell in worksheet[letter]:
            value = cell.value
            if value is None:
                continue
            value_length = len(str(value))
            if value_length > max_length:
                max_length = value_length
        worksheet.column_dimensions[letter].width = min(max(max_length + 2, 12), MAX_COLUMN_WIDTH)


def export_stores_to_excel(
    stores: list[StoreRecord],
    output_path: str = "output/stores.xlsx",
    snapshot_path: str | None = None,
    *,
    write_snapshot: bool = True,
    treat_diff_as_initial: bool = False,
) -> DiffResult:
    """Export stores, changes and statistics into an Excel workbook."""
    output_file = Path(output_path)
    output_file.parent.mkdir(parents=True, exist_ok=True)

    if snapshot_path is None:
        snapshot_file = output_file.with_name(f"{output_file.stem}_snapshot.json")
    else:
        snapshot_file = Path(snapshot_path)

    current_snapshot = build_snapshot_rows(stores)
    snapshot_load_result = load_snapshot_with_meta(snapshot_file)
    if snapshot_load_result.status == "invalid":
        logger.warning("Snapshot file is invalid and will be reinitialized: %s", snapshot_file)
    diff_result = compute_diff(
        previous_snapshot=snapshot_load_result.rows,
        current_snapshot=current_snapshot,
        treat_as_initial=treat_diff_as_initial or snapshot_load_result.status != "loaded",
        snapshot_status=snapshot_load_result.status,
    )

    data_df = _build_data_df(stores)
    changes_df = _build_changes_df(diff_result)
    stats_df = _build_stats_df(data_df)

    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        data_df.to_excel(writer, sheet_name=DATA_SHEET_NAME, index=False)
        changes_df.to_excel(writer, sheet_name=CHANGES_SHEET_NAME, index=False)
        stats_df.to_excel(writer, sheet_name=STATS_SHEET_NAME, index=False)
        for sheet_name in (DATA_SHEET_NAME, CHANGES_SHEET_NAME, STATS_SHEET_NAME):
            _apply_worksheet_formatting(writer.book[sheet_name])

    if write_snapshot:
        save_snapshot(stores, snapshot_file)
    else:
        logger.warning("Snapshot update skipped, baseline preserved: %s", snapshot_file)
    return diff_result
